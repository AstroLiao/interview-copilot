"""Interview Copilot 独立服务：静态页面 + AI代理 + 简历解析 + 手机推送中转 + 题库。

无数据库依赖：会话存浏览器 localStorage，推送中转存内存，题库存 bank.json。
运行：python server.py  （默认 http://0.0.0.0:8787）
"""
import io
import json
import os
import re
import socket
import time
import uuid
from pathlib import Path

import httpx
import uvicorn
from fastapi import FastAPI, File, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel

ROOT = Path(__file__).parent
STATIC = ROOT / "static"
BANK_FILE = ROOT / "bank.json"

# 页面存活追踪：每个标签页一个 ID，全部离线后服务自动退出
_tabs: dict = {}
_exitAt: float | None = None
_startedAt = time.time()

app = FastAPI(title="Interview Copilot", docs_url=None, redoc_url=None)


def _get_lan_ips() -> list[str]:
    hostname = socket.gethostname()
    candidates = set()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        candidates.add(s.getsockname()[0])
        s.close()
    except Exception:
        pass
    try:
        candidates.update(socket.gethostbyname_ex(hostname)[2])
    except Exception:
        pass
    skip = ("127.", "169.254.", "172.17.")
    return sorted(ip for ip in candidates if not ip.startswith(skip))


# ── 手机推送中转（内存态） ──
_relay = {"text": "", "ts": 0}


class RelayIn(BaseModel):
    text: str


@app.post("/api/relay")
async def push_relay(payload: RelayIn):
    text = payload.text.strip()
    if text:
        _relay["text"] = text
        _relay["ts"] += 1
    return {"ok": True, "ts": _relay["ts"]}


@app.get("/api/relay")
async def get_relay(after: int = 0):
    return {"text": _relay["text"], "ts": _relay["ts"], "lan_ips": _get_lan_ips()}


# ── AI 代理（绕开浏览器CORS；不兼容参数自动降级重试） ──
class AIChatIn(BaseModel):
    base: str
    key: str
    model: str
    messages: list
    temperature: float = 0.7
    max_tokens: int = 1500
    thinking_level: str = "low"


def _apply_thinking(payload: dict, level: str):
    if level in ("none", "minimal"):
        payload["thinking"] = {"type": "disabled"}
    elif level in ("low", "medium", "high", "xhigh", "max"):
        payload["reasoning_effort"] = level


@app.post("/api/ai/chat")
async def ai_chat(p: AIChatIn):
    base = p.base.strip().rstrip("/").removesuffix("/chat/completions")
    url = f"{base}/chat/completions"
    headers = {"Authorization": f"Bearer {p.key}", "Content-Type": "application/json"}
    payload = {
        "model": p.model,
        "messages": p.messages,
        "stream": True,
        "temperature": p.temperature,
        "max_tokens": p.max_tokens,
    }
    _apply_thinking(payload, p.thinking_level)
    clean_payload = {k: v for k, v in payload.items() if k not in ("thinking", "reasoning_effort")}

    async def gen():
        try:
            async with httpx.AsyncClient(timeout=httpx.Timeout(120, connect=15)) as client:
                body = payload
                for attempt in range(2):
                    async with client.stream("POST", url, headers=headers, json=body) as resp:
                        if resp.status_code == 400 and attempt == 0:
                            err_body = (await resp.aread()).decode("utf-8", "ignore")
                            if "param" in err_body or "invalid" in err_body.lower() or "unknown" in err_body.lower():
                                body = clean_payload
                                continue
                            err = json.dumps({"proxy_error": f"HTTP 400: {err_body[:300]}"}, ensure_ascii=False)
                            yield f"data: {err}\n\n"
                            return
                        if resp.status_code != 200:
                            eb = (await resp.aread()).decode("utf-8", "ignore")[:300]
                            err = json.dumps({"proxy_error": f"HTTP {resp.status_code}: {eb}"}, ensure_ascii=False)
                            yield f"data: {err}\n\n"
                            return
                        async for chunk in resp.aiter_bytes():
                            yield chunk
                        return
        except httpx.ConnectError as e:
            err = json.dumps({"proxy_error": f"无法连接到 {url}：{str(e)[:150]}（检查Base URL/网络）"}, ensure_ascii=False)
            yield f"data: {err}\n\n"
        except Exception as e:
            err = json.dumps({"proxy_error": f"代理转发异常: {str(e)[:150]}"}, ensure_ascii=False)
            yield f"data: {err}\n\n"

    return StreamingResponse(gen(), media_type="text/event-stream")


# ── 简历解析 ──
@app.post("/api/ai/resume")
async def parse_resume(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    data = await file.read()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        text += "\n" + " | ".join(cells)
        elif name.endswith(".doc"):
            return {"text": "", "error": "暂不支持老版 .doc 格式，请另存为 .docx 或 PDF 后重试"}
        else:
            text = data.decode("utf-8", "ignore")
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        return {"text": text[:20000], "filename": file.filename}
    except Exception as e:
        return {"text": "", "error": f"解析失败: {str(e)[:200]}"}


# ── 题库（bank.json 持久化） ──
def _load_bank() -> list:
    if BANK_FILE.exists():
        try:
            return json.loads(BANK_FILE.read_text(encoding="utf-8")).get("cards", [])
        except Exception:
            return []
    return []


def _save_bank(cards: list):
    BANK_FILE.write_text(json.dumps({"cards": cards}, ensure_ascii=False, indent=1), encoding="utf-8")


def _norm_q(s: str) -> str:
    return re.sub(r"[\s\W]+", "", s.lower())[:80]


@app.get("/api/bank")
async def bank_get():
    return {"cards": _load_bank()}


class BankImport(BaseModel):
    cards: list  # [{q, a, source?}]


@app.post("/api/bank/import")
async def bank_import(p: BankImport):
    cards = _load_bank()
    existing = {_norm_q(c["q"]): c for c in cards}
    added = replaced = skipped = 0
    for c in p.cards:
        q = (c.get("q") or "").strip()
        a = (c.get("a") or "").strip()
        if not q or not a:
            skipped += 1
            continue
        key = _norm_q(q)
        card = {
            "id": uuid.uuid4().hex[:8], "q": q, "a": a,
            "source": (c.get("source") or "导入")[:40],
            "createdAt": int(time.time() * 1000),
        }
        if key in existing:
            existing[key].update({"a": a, "source": card["source"]})
            replaced += 1
        else:
            cards.append(card)
            existing[key] = card
            added += 1
    _save_bank(cards)
    return {"ok": True, "added": added, "replaced": replaced, "skipped": skipped, "total": len(cards)}


class CardIn(BaseModel):
    q: str
    a: str


@app.post("/api/bank/card")
async def bank_add(c: CardIn):
    r = await bank_import(BankImport(cards=[{"q": c.q, "a": c.a, "source": "手动"}]))
    return r


@app.delete("/api/bank/card/{card_id}")
async def bank_delete(card_id: str):
    cards = _load_bank()
    cards = [c for c in cards if c["id"] != card_id]
    _save_bank(cards)
    return {"ok": True, "total": len(cards)}


@app.delete("/api/bank")
async def bank_clear():
    _save_bank([])
    return {"ok": True}


# ── 页面存活心跳：所有页面关闭后自动退出服务 ──
@app.post("/api/alive")
async def alive(request: Request):
    global _exitAt
    tab = None
    try:
        body = await request.body()
        if body:
            tab = json.loads(body).get("tab")
    except Exception:
        pass
    if tab:
        _tabs[tab] = time.time()
    _exitAt = None
    return {"ok": True, "tabs": len(_tabs)}


@app.post("/api/bye")
async def bye(request: Request):
    global _exitAt
    tab = None
    try:
        body = await request.body()
        if body:
            tab = json.loads(body).get("tab")
    except Exception:
        pass
    if tab and tab in _tabs:
        del _tabs[tab]
    if not _tabs:
        _exitAt = time.time() + 3  # 3 秒宽限：F5 刷新时新页面会立刻报活并取消
    return {"ok": True}


def _watchdog():
    global _exitAt
    while True:
        time.sleep(1)
        now = time.time()
        for t in [t for t, ts in _tabs.items() if now - ts > 90]:
            _tabs.pop(t, None)
        if _tabs and now - max(_tabs.values()) > 90:
            _tabs.clear()
        if _exitAt and now >= _exitAt and not _tabs:
            print("所有页面已关闭，服务自动退出")
            os._exit(0)
        if not _tabs and now - _startedAt > 90:
            print("启动后一直没有页面连接，服务自动退出")
            os._exit(0)


# ── 页面 ──
@app.get("/")
async def index():
    return FileResponse(STATIC / "index.html")


if __name__ == "__main__":
    import sys
    import threading
    import urllib.request
    import webbrowser

    URL = "http://localhost:8787"
    dev_mode = "--no-open" in sys.argv

    def already_up() -> bool:
        try:
            urllib.request.urlopen("http://127.0.0.1:8787/", timeout=1)
            return True
        except Exception:
            return False

    if already_up():
        print("服务已在运行，直接打开页面")
        webbrowser.open(URL)
        raise SystemExit(0)

    if not dev_mode:
        # 页面全部关闭后自动退出（--no-open 开发模式下不启用）
        threading.Thread(target=_watchdog, daemon=True).start()
        threading.Timer(1.5, lambda: webbrowser.open(URL)).start()

    uvicorn.run(app, host="0.0.0.0", port=8787, log_level="info")
